'''名单批次：从 CRM 拉取，或从 Excel / CSV 导入，落成可分配的名单条目。

导入一律走同一套校验：号码归一化 → 合法性 → 批内去重 → 库内去重 → 写库，并返回逐行报告，
让「哪几行没进来、为什么」在界面上说得清楚。
'''

from __future__ import annotations

import csv
import io
from typing import Any

import sqlalchemy as sa
from openpyxl import load_workbook

from app.adapters.sql import Store
from app.adapters.tables import crm_records, list_batches, list_items, task_items
from app.core.config import Settings
from app.core.ids import new_id
from app.core.phone import validate_phone
from app.domain import Actor
from app.errors import Conflict, Forbidden, NotFound, ValidationFailed
from app.integrations.crm import field_map

PHONE_HEADERS = ('phone', '手机号', '手机', '号码', '电话', '联系电话', '手机号码', '客户电话', '联系方式')
COMPANY_HEADERS = ('company', '公司', '公司名称', '企业', '客户名称')
CONTACT_HEADERS = ('contact', '联系人', '联系人姓名')
NAME_HEADERS = ('name', '客户', '客户名称', '公司', '公司名称', '企业', '姓名')

MAX_IMPORT_ROWS = 20000
MAX_CRM_PULL = 2000


class BatchService:
    def __init__(self, store: Store, settings: Settings, audit: Any, gateway: Any) -> None:
        self.store = store
        self.settings = settings
        self.audit = audit
        self.gateway = gateway

    # ---- 查询 ----
    def list_batches(
        self, actor: Actor, *, search: str | None = None, page: int = 1, page_size: int | None = None
    ) -> dict[str, Any]:
        size = self._page_size(page_size)
        offset = (max(1, int(page)) - 1) * size
        rows = self.store.select(
            list_batches,
            order_by='created_at',
            desc=True,
            limit=size,
            offset=offset,
            search=search,
            search_fields=('name', 'note', 'crm_object'),
        )
        total = self.store.count(list_batches)
        return {'data': rows, 'total': total, 'page': max(1, int(page)), 'page_size': size}

    def get_batch(self, actor: Actor, batch_id: str) -> dict:
        row = self.store.get(list_batches, batch_id)
        if row is None:
            raise NotFound('名单批次不存在')
        return row

    def list_items(
        self,
        actor: Actor,
        batch_id: str,
        *,
        status: str | None = None,
        search: str | None = None,
        page: int = 1,
        page_size: int | None = None,
    ) -> dict[str, Any]:
        batch = self.get_batch(actor, batch_id)
        size = self._page_size(page_size)
        offset = (max(1, int(page)) - 1) * size
        filters: dict[str, Any] = {'batch_id': batch_id}
        if status:
            filters['status'] = status
        rows = self.store.select(
            list_items,
            filters=filters,
            order_by='created_at',
            desc=False,
            limit=size,
            offset=offset,
            search=search,
            search_fields=('phone', 'name', 'company', 'contact_name'),
        )
        total = self.store.count(list_items, filters=filters)
        return {
            'data': [self.public_item(row, actor) for row in rows],
            'total': total,
            'batch': batch,
            'page': max(1, int(page)),
            'page_size': size,
        }

    def delete_batch(self, actor: Actor, batch_id: str) -> dict:
        if not actor.is_manager:
            raise Forbidden('需要主管及以上权限')
        batch = self.get_batch(actor, batch_id)
        item_ids = [row['id'] for row in self.store.select(list_items, filters={'batch_id': batch_id})]
        if item_ids and self.store.count(task_items, filters={'list_item_id': item_ids}) > 0:
            raise Conflict('该批次已被外呼任务引用，不能删除')
        self.store.delete_where(list_items, [list_items.c.batch_id == batch_id])
        self.store.delete(list_batches, batch_id)
        self.audit.log(actor, 'delete_batch', object_type='list_batches', object_id=batch_id)
        return {'detail': '已删除', 'id': batch_id, 'name': batch.get('name')}

    # ---- 从 CRM 拉取 ----
    def create_from_crm(self, actor: Actor, payload: dict) -> dict:
        if not actor.is_manager:
            raise Forbidden('需要主管及以上权限')
        crm_object = str(payload.get('crm_object') or 'leads')
        if not field_map.supports(crm_object):
            raise ValidationFailed('不支持从 CRM 拉取的对象：' + crm_object)
        limit = int(payload.get('limit') or 200)
        if limit < 1 or limit > MAX_CRM_PULL:
            raise ValidationFailed('单次拉取条数必须落在 1~' + str(MAX_CRM_PULL))
        search = payload.get('search')
        filters = payload.get('filters') or {}
        name = str(payload.get('name') or '').strip() or (
            'CRM ' + field_map.LIST_OBJECT_FIELDS[crm_object]['label'] + ' ' + crm_object
        )

        records, fetched = self._pull(crm_object, limit=limit, search=search, filters=filters)
        batch = self.store.insert(
            list_batches,
            {
                'id': new_id('batch'),
                'name': name,
                'source': 'crm',
                'crm_object': crm_object,
                'filter_json': {'search': search, 'filters': filters, 'limit': limit, 'fetched': fetched},
                'note': payload.get('note'),
                'created_by': actor.user_id,
            },
        )
        report = self._persist(actor, batch, records)
        self.audit.log(
            actor,
            'create_batch_from_crm',
            object_type='list_batches',
            object_id=batch['id'],
            detail={'crm_object': crm_object, 'fetched': fetched, 'imported': report['imported']},
        )
        return {'batch': self.store.get(list_batches, batch['id']), 'report': report}

    def sync_crm(
        self, actor: Actor, *, objects: list[str] | None = None, limit: int = 500
    ) -> dict[str, Any]:
        '''把 CRM 记录同步成本地快照，供工作台客户资料卡与黑名单客户维度使用。'''
        targets = [item for item in (objects or list(field_map.LIST_OBJECT_FIELDS)) if field_map.supports(item)]
        if not targets:
            raise ValidationFailed('没有可同步的 CRM 对象')
        summary: dict[str, Any] = {'objects': {}, 'total': 0, 'failed': 0}
        for crm_object in targets:
            try:
                records = self._fetch_raw(crm_object, limit=limit)
            except Exception as error:  # 上游不可达时只记失败，不让整批同步崩掉
                summary['objects'][crm_object] = {'ok': False, 'detail': str(error)}
                summary['failed'] += 1
                continue
            kept = 0
            for record in records:
                if self._upsert_snapshot(crm_object, record):
                    kept += 1
            summary['objects'][crm_object] = {'ok': True, 'synced': kept}
            summary['total'] += kept
        self.audit.log(actor, 'sync_crm', object_type='crm_records', detail={'total': summary['total']})
        return summary

    def get_snapshot(self, crm_object: str, record_id: str) -> dict | None:
        return self.store.find_one(
            crm_records, object_type=str(crm_object), record_id=str(record_id)
        )

    def ensure_snapshot(self, crm_object: str, record_id: str) -> dict | None:
        cached = self.get_snapshot(crm_object, record_id)
        if cached:
            return cached
        try:
            record = self.gateway.get_record(crm_object, record_id)
        except Exception:
            return None
        if not record:
            return None
        self._upsert_snapshot(crm_object, record)
        return self.get_snapshot(crm_object, record_id)

    # ---- 导入 ----
    def import_file(
        self, actor: Actor, *, filename: str, content: bytes, name: str | None = None, note: str | None = None
    ) -> dict:
        rows = self._parse_file(filename, content)
        if len(rows) > MAX_IMPORT_ROWS:
            raise ValidationFailed('单次导入不能超过 ' + str(MAX_IMPORT_ROWS) + ' 行')
        batch_name = str(name or '').strip() or (filename or '导入名单')
        records: list[dict[str, Any]] = []
        for index, row in enumerate(rows, start=1):
            records.append(
                {
                    'phone': row.get('phone'),
                    'name': row.get('name'),
                    'company': row.get('company'),
                    'contact_name': row.get('contact_name'),
                    'row': index,
                    'crm_object': None,
                    'crm_record_id': None,
                    'extra': {'source_file': filename},
                }
            )
        batch = self.store.insert(
            list_batches,
            {
                'id': new_id('batch'),
                'name': batch_name,
                'source': 'import',
                'crm_object': None,
                'filter_json': {'file': filename, 'rows': len(records)},
                'note': note,
                'created_by': actor.user_id,
            },
        )
        report = self._persist(actor, batch, records)
        self.audit.log(
            actor,
            'import_batch',
            object_type='list_batches',
            object_id=batch['id'],
            detail={'file': filename, 'imported': report['imported']},
        )
        return {'batch': self.store.get(list_batches, batch['id']), 'report': report}

    # ---- 内部 ----
    def _persist(self, actor: Actor, batch: dict, records: list[dict]) -> dict[str, Any]:
        existing = self._existing_phones()
        seen: set[str] = set()
        valid: list[dict] = []
        duplicates: list[dict] = []
        failed: list[dict] = []
        for record in records:
            raw_phone = record.get('phone')
            ok, phone, reason = validate_phone(raw_phone or '')
            if not ok:
                failed.append({'row': record.get('row'), 'phone': str(raw_phone or ''), 'reason': reason})
                continue
            if phone in seen or phone in existing:
                duplicates.append({'row': record.get('row'), 'phone': phone, 'reason': '号码重复'})
                continue
            seen.add(phone)
            valid.append(
                {
                    'id': new_id('item'),
                    'batch_id': batch['id'],
                    'phone': phone,
                    'name': _clean(record.get('name')),
                    'company': _clean(record.get('company')),
                    'contact_name': _clean(record.get('contact_name')),
                    'crm_object': record.get('crm_object') or None,
                    'crm_record_id': record.get('crm_record_id') or None,
                    'status': 'available',
                    'attempts': 0,
                    'extra': record.get('extra') or {},
                }
            )
        if valid:
            self.store.insert_many(list_items, valid)
        report = {
            'total_rows': len(records),
            'imported': len(valid),
            'duplicates': len(duplicates),
            'failed': len(failed),
            'duplicate_rows': duplicates[:50],
            'failed_rows': failed[:50],
        }
        self.store.update(
            list_batches,
            batch['id'],
            {
                'total': len(records),
                'imported_total': len(valid),
                'duplicate_total': len(duplicates),
                'failed_total': len(failed),
            },
        )
        return report

    def _fetch_raw(
        self, crm_object: str, *, limit: int, search: str | None = None, filters: dict | None = None
    ) -> list[dict]:
        '''按页拉原始 CRM 记录（未拍平），上限 limit 条。'''
        collected: list[dict] = []
        cursor: str | None = None
        page_size = min(int(self.settings.crm_page_size or 100), 200)
        while len(collected) < limit:
            page = self.gateway.list_records(
                crm_object, limit=page_size, cursor=cursor, search=search, filters=filters
            )
            rows = list(page.get('data') or [])
            for record in rows:
                collected.append(record)
                if len(collected) >= limit:
                    break
            cursor = page.get('next_cursor')
            if not rows or not cursor:
                break
        return collected

    def _pull(
        self, crm_object: str, *, limit: int, search: str | None = None, filters: dict | None = None
    ) -> tuple[list[dict], int]:
        '''返回 (拍平后的名单条目, 上游返回条数)；顺带把原始记录落成本地快照。'''
        records = self._fetch_raw(crm_object, limit=limit, search=search, filters=filters)
        mapped: list[dict] = []
        for index, record in enumerate(records, start=1):
            self._upsert_snapshot(crm_object, record)
            item = field_map.map_record(crm_object, record)
            item['row'] = index
            mapped.append(item)
        return mapped, len(records)

    def _upsert_snapshot(self, crm_object: str, record: dict) -> bool:
        record_id = str(record.get('id') or '').strip()
        if not record_id:
            return False
        values = {
            'title': field_map.title_of(crm_object, record),
            'phone': str(record.get('phone') or '').strip() or None,
            'name': str(record.get('name') or record.get('company') or record.get('contact') or '').strip() or None,
            'company': str(record.get('company') or '').strip() or None,
            'owner_id': record.get('owner_id'),
            'team_id': record.get('team_id'),
            'stage': record.get('stage'),
            'extra': {key: value for key, value in record.items() if key not in ('id',)},
            'synced_at': _now(),
            'deleted': False,
        }
        existing = self.store.find_one(crm_records, object_type=str(crm_object), record_id=record_id)
        conditions = [
            crm_records.c.object_type == str(crm_object),
            crm_records.c.record_id == record_id,
        ]
        if existing:
            self.store.update_where(crm_records, conditions, values)
        else:
            self.store.insert(
                crm_records, {'object_type': str(crm_object), 'record_id': record_id, **values}
            )
        return True

    def _existing_phones(self) -> set[str]:
        rows = self.store.raw(sa.select(list_items.c.phone).distinct())
        return {str(row['phone']) for row in rows if row.get('phone')}

    def _parse_file(self, filename: str, content: bytes) -> list[dict[str, Any]]:
        lowered = str(filename or '').lower()
        if lowered.endswith('.xlsx') or lowered.endswith('.xlsm'):
            return self._parse_xlsx(content)
        if lowered.endswith('.csv') or lowered.endswith('.txt'):
            return self._parse_csv(content)
        raise ValidationFailed('只支持 .csv / .xlsx 名单文件，收到：' + str(filename))

    def _parse_csv(self, content: bytes) -> list[dict[str, Any]]:
        text = _decode(content)
        reader = csv.reader(io.StringIO(text))
        try:
            header = next(reader)
        except StopIteration:
            raise ValidationFailed('文件是空的')
        mapping = _header_map(header)
        out: list[dict[str, Any]] = []
        for row in reader:
            if not any(str(cell or '').strip() for cell in row):
                continue
            out.append(_row_from(mapping, row))
        if mapping.get('phone') is None:
            raise ValidationFailed('没找到号码列，请把表头命名为 phone / 手机号 / 电话')
        return out

    def _parse_xlsx(self, content: bytes) -> list[dict[str, Any]]:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            header = [str(cell) if cell is not None else '' for cell in next(rows)]
        except StopIteration:
            raise ValidationFailed('工作表是空的')
        mapping = _header_map(header)
        if mapping.get('phone') is None:
            raise ValidationFailed('没找到号码列，请把表头命名为 phone / 手机号 / 电话')
        out: list[dict[str, Any]] = []
        for row in rows:
            values = ['' if cell is None else cell for cell in row]
            if not any(str(cell).strip() for cell in values):
                continue
            out.append(_row_from(mapping, values))
        return out

    def _page_size(self, value: int | None) -> int:
        size = int(value or self.settings.default_page_size)
        return max(1, min(size, int(self.settings.max_page_size)))

    @staticmethod
    def public_item(row: dict, actor: Actor) -> dict:
        '''名单列表一律脱敏；看全号只有通话详情一条路（管理员，且写审计）。'''
        from app.core.phone import mask_phone

        item = dict(row)
        item['phone_masked'] = mask_phone(str(row.get('phone') or ''))
        item['phone'] = item['phone_masked']
        return item


def _row_from(mapping: dict[str, int], row: Any) -> dict[str, Any]:
    def cell(key: str) -> str:
        index = mapping.get(key)
        if index is None or index >= len(row):
            return ''
        return str(row[index] if row[index] is not None else '').strip()

    return {
        'phone': cell('phone'),
        'name': cell('name'),
        'company': cell('company'),
        'contact_name': cell('contact_name'),
    }


def _header_map(header: list[Any]) -> dict[str, int]:
    normalized = {str(cell or '').strip().lower(): index for index, cell in enumerate(header)}
    mapping: dict[str, int] = {}
    for key, candidates in (
        ('phone', PHONE_HEADERS),
        ('company', COMPANY_HEADERS),
        ('contact_name', CONTACT_HEADERS),
        ('name', NAME_HEADERS),
    ):
        for candidate in candidates:
            if candidate in normalized:
                mapping[key] = normalized[candidate]
                break
    return mapping


def _decode(content: bytes) -> str:
    for encoding in ('utf-8-sig', 'utf-8', 'gbk', 'gb18030'):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode('utf-8', errors='replace')


def _clean(value: Any) -> str | None:
    text = str(value or '').strip()
    return text or None


def _now() -> str:
    from app.core.clock import now_iso

    return now_iso()